"""Fonction pour la création des différents graphiques"""

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.chart import (
    ScatterChart,
    BubbleChart,
    Reference,
    BarChart,
    RadarChart,
    #existe bien
    Series
)

def scatter_chart( 
        worksheet_chart: Worksheet, 
        worksheet_data: Worksheet, 
        where: str, 
        col_x: int, 
        col_y: int, 
        min_row: int, 
        max_row: int
        ) -> None :
    """Création d'un graphique nuage de points

    Args:
        worksheet_chart (Worksheet): Worksheet où ajouter le graphique
        worksheet_data (Worksheet): Worksheet où prendre les données
        where (str): Cellule où ajouter le graphique
        col_x (int): colonne des données d'abscisse
        col_y (int): colonne des données d'ordonnée
        min_row (int): début des données
        max_row (int): fin des données
    """
    
    c1 = ScatterChart()
    c1.title = "Complexité vs Note utilisateur"
    c1.legend = None
    c1.y_axis.title = 'Note utilisateur'
    c1.x_axis.title = 'Complexité'
    c1.height = 12

    x_values = Reference(worksheet_data, min_col=col_x, min_row=min_row, max_row=max_row)
    y_values = Reference(worksheet_data, min_col=col_y, min_row=min_row, max_row=max_row)
    series = Series(y_values, x_values, title_from_data=False)

    series.marker.symbol = "circle" 
    series.marker.size = 3
    series.graphicalProperties.line.noFill = True
    
    c1.series.append(series)

    worksheet_chart.add_chart(c1, where)


def bar_chart( 
        worksheet_chart: Worksheet, 
        worksheet_data: Worksheet, 
        where: str, 
        col_x: int, 
        col_y: int, 
        min_row: int, 
        max_row: int
        ) -> None :
    """Création d'un graphique en bar

    Args:
        worksheet_chart (Worksheet): Worksheet où ajouter le graphique
        worksheet_data (Worksheet): Worksheet où prendre les données
        where (str): Cellule où ajouter le graphique
        col_x (int): colonne des données d'abscisse
        col_y (int): colonne des données d'ordonnée
        min_row (int): début des données
        max_row (int): fin des données
    """
    
    c1 = BarChart()
    c1.title="Top 10 BGG : Nombre d'unités de jeux possédées"
    c1.type="col"
    c1.legend = None
    c1.y_axis.title = "Nombre d'unités jeux possédés"
    c1.x_axis.title = 'Jeux'
    c1.height = 12

    x_values = Reference(worksheet_data, min_col=col_x, min_row=min_row, max_row=max_row)
    y_values = Reference(worksheet_data, min_col=col_y, min_row=min_row, max_row=max_row)

    c1.add_data(y_values, titles_from_data=False)
    c1.set_categories(x_values)

    worksheet_chart.add_chart(c1, where)


def radar_chart( 
        worksheet_chart: Worksheet, 
        worksheet_data: Worksheet, 
        where: str, 
        min_row: int,
        max_row: int, 
        min_col: int 
        ) -> None :
    """Création d'un graphique radar

    Args:
        worksheet_chart (Worksheet): Worksheet où ajouter le graphique
        worksheet_data (Worksheet): Worksheet où prendre les données
        where (str): Cellule où ajouter le graphique
        min_row (int): début des données
        max_row (int): fin des données
        min_col (int): colonne des labels
    """
    c1=RadarChart()
    c1.title = "Carte d'identité de la famille de jeux"
    labels=Reference(worksheet_data, min_col=min_col,min_row=min_row+1,max_row=max_row)
    y_values=Reference(worksheet_data, min_col=min_col+1,max_col=min_col+2,min_row=min_row,max_row=max_row)
    c1.add_data(y_values, titles_from_data=True)
    c1.set_categories(labels)
    # c1.y_axis.delete=True
    c1.height = 12

    worksheet_chart.add_chart(c1, where)